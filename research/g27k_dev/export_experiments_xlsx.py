#!/usr/bin/env python3
"""The extreme-risk experiments (backtest only, never for live use) as one
Excel workbook: gold_moonshot, max_growth, gold_sweep_cap, gold_week.

Every cell is a computed value (no formulas): the session's LibreOffice could
not recalculate, and the workbook is meant to open on a phone, where an
uncalculated formula shows blank. Dollar columns use $100 per account.

Usage: python3 research/g27k_dev/export_experiments_xlsx.py --out <xlsx>
"""
import argparse
import json
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
F = "Arial"
HEAD = PatternFill("solid", fgColor="E8E2D0")
INPUT = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
SIG_TH = {"every": "ทะลุ High/Low 20 แท่ง", "trend": "ใกล้ High/Low 55 แท่ง + D1 ทางเดียวกัน",
          "m30rule": "กฎ M30: trend + ATR14/ATR100 ≥ 1.5", "valid": "รูปแบบ M15 ที่พิสูจน์แล้ว: ช่วงบน 10% ของ 250 แท่ง + ATR ≥ 1.5 + D1"}
RULE_TH = {"A": "M15 พิสูจน์แล้ว: pos250≥0.9, ATR≥1.5, D1 ทางเดียวกัน, ออก channel 20",
           "B": "M15 พิสูจน์แล้ว: pos250≥0.9, ATR≥1.25, D1 ทางเดียวกัน, ออก channel 20",
           "C": "กฎ M30 ใช้บน M15 (ยังไม่พิสูจน์บน M15), TP 2R"}


def sheet(wb, title, headers, widths, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=F, italic=True, color="595959")
        r0 = 3
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(r0, j, h)
        c.font = Font(name=F, bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 44
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws, r0 + 1


def put(ws, r, values, fmts):
    for j, (v, f) in enumerate(zip(values, fmts), 1):
        c = ws.cell(r, j, v)
        c.font = Font(name=F)
        c.border = THIN
        if f:
            c.number_format = f
    return r + 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    wb = Workbook()
    rd = wb.active
    rd.title = "อ่านก่อน"
    rd.column_dimensions["A"].width = 30
    rd.column_dimensions["B"].width = 100
    rows = [("การทดลองความเสี่ยงสูงสุด (backtest เท่านั้น ไม่ใช้เทรดจริง)", None),
            (None, None),
            ("เงินต้นต่อพอร์ตที่ใช้ในคอลัมน์ $", 100),
            (None, None),
            ("ชีต", "เนื้อหา"),
            ("Moonshot 3 เดือน", "ทองเสี่ยง 100% ต่อไม้ TP 5/7.5/10R บน M5 และ M15 ช่วง ก.ค.–ก.ย. 2026 เทียบกับตั้งค่าเดียวกัน 15 ปี"),
            ("Max growth", "ความเสี่ยงต่อไม้ที่ทำให้ G27K-F + M30 + H1 โตสูงสุดใน 15 ปี และผลเมื่อปรับขึ้นลง ×0.25 ถึง ×4"),
            ("Max growth grid", "ทุกชุดที่ลอง 810 ชุดต่อบัญชี (G27K-F 1–15%, M30 0–10%, H1 0–10%)"),
            ("ทองถอนกำไร รายเดือน", "ทอง M30 และ M15 เสี่ยง 5–100% ถอนกำไรทันที พอร์ตแตกแล้วเติมใหม่ เงื่อนไขแตกไม่เกิน 1 ครั้งต่อ 3 เดือน"),
            ("ทอง M15 รายสัปดาห์", "ทอง M15 สามกฎ เสี่ยง 5–100% ผลสุทธิต่อสัปดาห์"),
            (None, None),
            ("ตั้งค่าร่วม", "ราคาทอง Dukascopy M1 (กลาง bid/ask) · เข้าราคาเปิดแท่งถัดไป · SL 2 × ATR20 ของ TF นั้น · เดิน SL/TP บน M1 (ชน SL ก่อนถ้าโดนทั้งสองในนาทีเดียว)"),
            ("ต้นทุน", "spread 2 bp ของราคาต่อรอบ + swap ต่อคืน ตามสเปกที่ใช้ใน G27K · ไม่ได้คิด slippage"),
            ("ถอนกำไร / พอร์ตแตก", "ยอดเกินเงินต้นถอนออกทันที · เหลือต่ำกว่า 10% ของต้นนับเป็นพอร์ตแตกและเติมเงินต้นใหม่ทันที"),
            ("เงินติดลบ", "สมมติว่าโบรกมี negative balance protection"),
            ("Max growth", "บัญชีเดียวทบต้น ไม่มีถอน ความเสี่ยง % ของ balance ตอนเข้าไม้ ไม่มีเพดานไม้เปิดซ้อน · เงินสุดท้ายหลายล้านล้านเท่าเป็นตัวเลขของโมเดล ในตลาดจริงเปิดไม้ขนาดนั้นไม่ได้"),
            ("ตัวเลข", "ทุกช่องเป็นค่าที่คำนวณแล้ว (ไม่มีสูตร เปิดบนมือถือได้) · คอลัมน์ $ คิดที่เงินต้น $100 ต่อพอร์ต · ถ้าเงินต้นต่างไป ให้คูณคอลัมน์ 'เท่าของต้น' ด้วยเงินต้นนั้น"),
            ("ที่มา", "research/g27k_dev/gold_moonshot.py, max_growth.py, gold_sweep_cap.py, gold_week.py และไฟล์ .json ของแต่ละตัว")]
    for i, (k, v) in enumerate(rows, 1):
        c1, c2 = rd.cell(i, 1, k), rd.cell(i, 2, v)
        c1.font = Font(name=F, bold=i in (1, 5) or bool(k and v and i > 11))
        c2.font = Font(name=F, bold=i == 5)
        c2.alignment = Alignment(wrap_text=True, vertical="top")
    rd["A1"].font = Font(name=F, bold=True, size=14)
    rd["B3"].number_format = "$#,##0"
    rd["B3"].comment = Comment("เงินต้นต่อพอร์ตที่ใช้คิดคอลัมน์ $ ในชีตอื่น ($100 ตามตัวอย่างในแชต)", "Claude")
    CAPV = 100.0                     # dollars per account in the $ columns

    # Moonshot
    m = json.loads((HERE / "gold_moonshot.json").read_text())
    hdr = ["TF", "สัญญาณ", "คำอธิบาย", "TP (R)", "กำไรต่อไม้ชนะ", "ไม้ 3 เดือน", "ไม้ต่อวัน", "ชนะ", "ชนะขั้นต่ำที่เสมอตัว",
           "พอร์ตแตก", "เติมเงิน (เท่าของต้น)", "ถอนได้ (เท่าของต้น)", "คงเหลือ (เท่าของต้น)", "สุทธิ (เท่าของต้น)", "สุทธิ ($)", "ไม้ดีสุด",
           "ไม้ 15 ปี", "ชนะ 15 ปี", "สุทธิ 15 ปี (เท่าของต้น)", "สุทธิต่อไม้ 15 ปี"]
    ws, r = sheet(wb, "Moonshot 3 เดือน", hdr, [6, 10, 44, 7, 10, 9, 8, 8, 11, 9, 11, 11, 11, 11, 11, 9, 9, 9, 12, 11],
                  "ทองเสี่ยง 100% ต่อไม้ · ช่วง 1 ก.ค. – 30 ก.ย. 2026 · เรียงตามจำนวนไม้ 3 เดือน")
    items = sorted(m.items(), key=lambda kv: -kv[1]["last3m"]["trades"])
    for key, v in items:
        tf, sig, tp = key.split()
        l3, fy = v["last3m"], v["15y"]
        bal = l3["net"] - l3["withdrawn"] + l3["deposited"]
        k_ = float(tp[:-1])
        put(ws, r, [tf, sig, SIG_TH[sig], k_, k_, l3["trades"], l3["trades"] / 92, l3["win"], 1 / (k_ + 1), l3["blowups"],
                    l3["deposited"], l3["withdrawn"], bal, l3["net"], l3["net"] * CAPV, l3["best_trade"], fy["trades"], fy["win"], fy["net"],
                    fy["net"] / fy["trades"] if fy["trades"] else 0.0],
            [None, None, None, "0.0", "0%", "#,##0", "0.0", "0.0%", "0.0%", "#,##0", "0.0", "0.0", "0.00", "+0.0;-0.0;0.0", "$#,##0;($#,##0);-", "0%",
             "#,##0", "0.0%", "+0;-0;0", "+0.00;-0.00;0.00"])
        r += 1
    ws.cell(r + 1, 1, "ไม้ต่อวัน = ไม้ ÷ 92 วันของช่วง ก.ค.–ก.ย. · ชนะขั้นต่ำที่เสมอตัว = 1 ÷ (TP + 1) · กำไรต่อไม้ชนะ = TP × 100% ก่อนหักต้นทุน").font = Font(name=F, italic=True, color="595959")

    # Max growth (scaled)
    g = json.loads((HERE / "max_growth.json").read_text())
    hdr = ["บัญชี", "ตัวคูณ", "G27K-F เสี่ยงต่อไม้", "M30 เสี่ยงต่อไม้", "H1 เสี่ยงต่อไม้", "เงินสุดท้าย (เท่าของต้น)", "ต่อปี", "DD สูงสุด",
           "3 เดือนแย่สุด", "เดือนแย่สุด", "DD เกิน 50% ครั้งแรก", "พอร์ตแตก", "ต่อปี 2011–18", "DD 2011–18", "ต่อปี 2019–26", "DD 2019–26", "เงินสุดท้าย ($)"]
    ws, r = sheet(wb, "Max growth", hdr, [10, 8, 11, 11, 11, 15, 9, 9, 10, 10, 14, 12, 10, 10, 10, 10, 18],
                  "จุดโตสูงสุดหาจาก grid 810 ชุดต่อบัญชี (ตัวคูณ 1) แล้วปรับทุกค่าพร้อมกัน · ทบต้น ไม่มีถอน · เริ่ม ก.ย. 2011")
    for acct in ("Cent", "Standard"):
        best = g[acct]["best"]
        base = r
        ws.cell(r, 1, f"{acct}: จุดสูงสุด G27K-F {best['kg']}% / M30 {best['km']}% / H1 {best['kh']}%").font = Font(name=F, bold=True)
        r += 1
        first = r
        for s in g[acct]["scaled"]:
            f_, h1, h2 = s["full"], s["h1"], s["h2"]
            put(ws, r, [acct, s["lam"], s["lam"] * best["kg"] / 100, s["lam"] * best["km"] / 100, s["lam"] * best["kh"] / 100, max(f_["final"], 0.0),
                        f_["cagr"], f_["dd"], f_.get("worst_3m"), f_.get("worst_month"), f_.get("first_dd50") or "-", f_.get("ruin_date") or "-",
                        h1["cagr"], h1["dd"], h2["cagr"], h2["dd"], max(f_["final"], 0.0) * CAPV],
                [None, "0.00\"×\"", "0.00%", "0.00%", "0.00%", "0.00E+00", "0%", "0%", "0%", "0%", None, None, "0%", "0%", "0%", "0%", "$0.00E+00"])
            r += 1
        r += 1
    ws.cell(r, 1, "เสี่ยงต่อไม้ = ตัวคูณ × ค่าที่จุดสูงสุด · เงินสุดท้ายติดลบแสดงเป็น 0 (พอร์ตแตก)").font = Font(name=F, italic=True, color="595959")

    # Max growth grid
    hdr = ["บัญชี", "G27K-F เสี่ยงต่อไม้", "M30 เสี่ยงต่อไม้", "H1 เสี่ยงต่อไม้", "เงินสุดท้าย (เท่าของต้น)", "ต่อปี", "DD สูงสุด", "เดือนแย่สุด", "3 เดือนแย่สุด",
           "ไม้", "DD เกิน 50% ครั้งแรก", "พอร์ตแตก"]
    ws, r = sheet(wb, "Max growth grid", hdr, [10, 11, 11, 11, 15, 9, 9, 10, 10, 9, 14, 12])
    for acct in ("Cent", "Standard"):
        for x in sorted(g[acct]["grid"], key=lambda x: -x["final"]):
            put(ws, r, [acct, x["kg"] / 100, x["km"] / 100, x["kh"] / 100, max(x["final"], 0.0), x["cagr"], x["dd"], x.get("worst_month"), x.get("worst_3m"),
                        x["n"], x.get("first_dd50") or "-", x.get("ruin_date") or "-"],
                [None, "0%", "0%", "0%", "0.00E+00", "0%", "0%", "0%", "0%", "#,##0", None, None])
            r += 1
    ws.auto_filter.ref = f"A1:L{r - 1}"

    # gold sweep with blow-up cap (monthly)
    c = json.loads((HERE / "gold_sweep_cap.json").read_text())
    hdr = ["TF", "เสี่ยงต่อไม้", "พอร์ตแตก 15 ปี", "แตกมากสุดใน 3 เดือน", "ผ่านเงื่อนไข (≤1 ครั้ง/3 เดือน)", "สุทธิเฉลี่ย/เดือน", "สุทธิกลาง/เดือน",
           "เดือนที่ได้ ≥ +5%", "เดือนดีสุด", "เดือนแย่สุด", "สุทธิ 15 ปี (เท่าของต้น)", "สุทธิ 15 ปี ($)", "สุทธิเฉลี่ย/เดือน ($)"]
    ws, r = sheet(wb, "ทองถอนกำไร รายเดือน", hdr, [6, 10, 11, 12, 14, 12, 12, 12, 10, 10, 13, 13, 13],
                  "ทองเท่านั้น · M30 = กฎที่พิสูจน์แล้ว 710 ไม้ · M15 = กฎเดียวกันบน M15 (ยังไม่พิสูจน์) · สุทธิ = ถอนได้ − เงินที่เติมหลังพอร์ตแตก เทียบเงินต้น")
    for tf in ("M30", "M15"):
        for k, x in c[tf]["risks"].items():
            put(ws, r, [tf, float(k), x["blowups"], x["worst_3m_blowups"], "ผ่าน" if x["ok"] else "ไม่ผ่าน", x["month_mean"], x["month_median"], x["months_ge5"],
                        x["best_month"], x["worst_month"], x["net_total"], x["net_total"] * CAPV, x["month_mean"] * CAPV],
                [None, "0%", "#,##0", "0", None, "+0.0%;-0.0%;0.0%", "+0.0%;-0.0%;0.0%", "0%", "+0%;-0%;0%", "+0%;-0%;0%", "+0.0;-0.0;0.0",
                 "$#,##0;($#,##0);-", "$#,##0.00;($#,##0.00);-"])
            r += 1
        r += 1

    # gold M15 weekly
    w = json.loads((HERE / "gold_week.json").read_text())
    hdr = ["กฎ", "คำอธิบาย", "ไม้ 15 ปี", "R เฉลี่ยต่อไม้", "เสี่ยงต่อไม้", "สุทธิเฉลี่ย/สัปดาห์", "สุทธิกลาง/สัปดาห์", "สัปดาห์ที่ได้ ≥ +5%",
           "สัปดาห์ติดลบ", "สัปดาห์ดีสุด", "พอร์ตแตก 15 ปี", "แตกต่อปี", "แตกมากสุดใน 3 เดือน", "สุทธิ 15 ปี (เท่าของต้น)", "สุทธิเฉลี่ย/สัปดาห์ ($)"]
    ws, r = sheet(wb, "ทอง M15 รายสัปดาห์", hdr, [5, 50, 9, 10, 9, 12, 12, 12, 10, 10, 11, 9, 12, 13, 13],
                  "ทอง M15 · ถอนกำไรทันที · สุทธิ = ถอนได้ − เงินที่เติมหลังพอร์ตแตก เทียบเงินต้น · 2011-09 ถึง 2026-09")
    for k, v in w.items():
        for rk, x in v["risks"].items():
            put(ws, r, [k, RULE_TH[k], v["n"], v["mean_R"], float(rk), x["week_mean"], x["week_median"], x["weeks_ge5"], x["weeks_neg"], x["best_week"],
                        x["blowups"], x["per_year"], x["worst_3m"], x["net_total"], x["week_mean"] * CAPV],
                [None, None, "#,##0", "+0.000;-0.000", "0%", "+0.0%;-0.0%;0.0%", "+0.0%;-0.0%;0.0%", "0%", "0%", "+0%;-0%;0%", "#,##0", "0.0", "0",
                 "+0.0;-0.0;0.0", "$#,##0.00;($#,##0.00);-"])
            r += 1
        r += 1
    wb.save(a.out)
    print("written", a.out)


if __name__ == "__main__":
    main()
